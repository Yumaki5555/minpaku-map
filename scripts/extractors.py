"""行政サイトからダウンロードした元ファイル(pdf/xlsx/xls/csv/docx)を
表(行のリスト)に変換するための共通処理。

自治体ごとに列の並びがバラバラなため、ここでは「何行あるか」
「住所らしき列・施設名らしき列はどれか」をヒューリスティックに
判定するところまでを共通化し、区ごとの細かい列マッピングは
後続のステップ(整形)で必要に応じて上書きする想定。
"""
from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field

import pandas as pd

# 一部の行政PDFは、フォントの不具合で「民」「亀」などの漢字が
# 見た目は同じだが別のUnicode文字(部首用の記号)に置き換わって
# 抽出されることがある。NFKC正規化で大半は直るが、この2文字は
# 正規化でも変換されないため、個別に置き換える。
_RADICAL_OVERRIDES = {
    "⺠": "民",  # CJK RADICAL CIVILIAN → 民
    "⻲": "亀",  # CJK RADICAL J-SIMPLIFIED TURTLE → 亀
}


def _fix_pdf_text(value):
    if not isinstance(value, str):
        return value
    text = unicodedata.normalize("NFKC", value)
    for bad, good in _RADICAL_OVERRIDES.items():
        text = text.replace(bad, good)
    return text

ADDRESS_HINTS = ["住所", "所在地", "所在", "location", "address"]
NAME_HINTS = ["施設名", "名称", "屋号", "建物名", "施設の名称", "住宅宿泊事業者", "営業者", "name"]

# 複数の表を1枚のPDFに載せている資料では、表が切り替わるたびに
# 見出し行(列名)がもう一度データとして紛れ込むことがある。
# セルの中身が丸ごとこれらの見出し語と一致する行は、見出し行とみなして除外する。
HEADER_LIKE_TOKENS = {
    "施設名称", "施設名", "名称", "所在地", "住所", "営業種別", "業種別", "種別", "区分",
    "施設種別", "対応する外国語", "no", "no.", "届出番号", "届出年月日", "許可番号", "許可日",
    "登録番号", "連絡先", "商号又は名称", "住宅宿泊事業者", "営業者", "建物名・部屋番号",
    "建物名", "部屋番号", "電話番号", "tel", "tel1", "備考", "住宅の住所",
}


def _is_header_like_row(row: list) -> bool:
    cells = [str(c).strip() for c in row if c not in (None, "")]
    if not cells:
        return False
    return all(c.lower() in HEADER_LIKE_TOKENS for c in cells)

_ZEN2HAN = str.maketrans("０１２３４５６７８９", "0123456789")

_ANCHOR = r"\s*(?:現在|時点|時点で)"
_REIWA_ANCHORED = [
    re.compile(r"令和\s*(\d{1,2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日" + _ANCHOR),
    re.compile(r"[Rr]\s*(\d{1,2})[.\s](\d{1,2})[.\s](\d{1,2})" + _ANCHOR),
]
_REIWA_PATTERNS = [
    re.compile(r"令和\s*(\d{1,2})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日"),
    re.compile(r"[Rr]\s*(\d{1,2})[.\s](\d{1,2})[.\s](\d{1,2})"),
]
_SEIREKI_ANCHORED = re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日" + _ANCHOR)
_SEIREKI_PATTERN = re.compile(r"(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日")


def _reiwa_to_str(m: "re.Match") -> str | None:
    r, mo, d = (int(x) for x in m.groups())
    year = r + 2018
    if 1 <= mo <= 12 and 1 <= d <= 31:
        return f"{year:04d}/{mo:02d}/{d:02d}"
    return None


def _seireki_to_str(m: "re.Match") -> str | None:
    y, mo, d = (int(x) for x in m.groups())
    if 2000 <= y <= 2100 and 1 <= mo <= 12 and 1 <= d <= 31:
        return f"{y:04d}/{mo:02d}/{d:02d}"
    return None


def find_source_date(text: str | None, allow_unanchored: bool = True) -> str | None:
    """テキストの中から「令和8年8月31日現在」のような日付表記を探し、
    見つかれば "2026/08/31" の形式で返す。見つからなければ None。

    「現在」「時点」が直後に続く日付(=一覧の基準日である可能性が高い)を
    優先する。allow_unanchored=True のときのみ、見つからなかった場合に
    それ以外の日付表記にもフォールバックする。

    表の本文(届出年月日などのデータ列)を渡す場合は allow_unanchored=False
    にすること。そうしないと、無関係なデータ行の日付を「基準日」だと
    誤認識してしまう(タイトルに「現在」の記載がない書類で発生しやすい)。
    """
    if not text:
        return None
    text = text.translate(_ZEN2HAN)

    for pat in _REIWA_ANCHORED:
        m = pat.search(text)
        if m:
            result = _reiwa_to_str(m)
            if result:
                return result
    m = _SEIREKI_ANCHORED.search(text)
    if m:
        result = _seireki_to_str(m)
        if result:
            return result

    if not allow_unanchored:
        return None

    for pat in _REIWA_PATTERNS:
        m = pat.search(text)
        if m:
            result = _reiwa_to_str(m)
            if result:
                return result
    m = _SEIREKI_PATTERN.search(text)
    if m:
        return _seireki_to_str(m)
    return None


@dataclass
class ExtractResult:
    row_count: int
    header: list[str] = field(default_factory=list)
    rows: list[list] = field(default_factory=list)
    sheet_name: str | None = None
    note: str = ""
    raw_text: str = ""


def _pick_header_row(raw_rows: list[list]) -> int:
    """先頭付近の行から、住所/施設名らしき語を含む行をヘッダー行とみなす。
    見つからなければ 0 行目(先頭行)をヘッダーとみなす。
    """
    limit = min(10, len(raw_rows))
    for i in range(limit):
        cells = [str(c).strip() for c in raw_rows[i] if c is not None]
        joined = "".join(cells)
        if any(h in joined for h in ADDRESS_HINTS) and any(h in joined for h in NAME_HINTS):
            return i
    for i in range(limit):
        cells = [str(c).strip() for c in raw_rows[i] if c is not None]
        joined = "".join(cells)
        if any(h in joined for h in ADDRESS_HINTS):
            return i
    return 0


def _clean_table(raw_rows: list[list]) -> ExtractResult:
    raw_rows = [r for r in raw_rows if any(c not in (None, "") for c in r)]
    if not raw_rows:
        return ExtractResult(row_count=0, note="表が空でした")
    header_idx = _pick_header_row(raw_rows)
    header = [str(c).strip() if c is not None else "" for c in raw_rows[header_idx]]
    header_key = "|".join(header)
    data_rows = raw_rows[header_idx + 1:]
    data_rows = [r for r in data_rows if any(c not in (None, "") for c in r)]
    # 複数ページのPDFでは見出し行がページごとに繰り返されることがあるため、
    # ヘッダーと完全一致する行はデータ行から除外する
    data_rows = [
        r for r in data_rows
        if "|".join(str(c).strip() if c is not None else "" for c in r) != header_key
        and not _is_header_like_row(r)
    ]
    # 日付の手がかり探索用に、表の先頭(タイトル行・見出し行)のテキストだけを残す。
    # ヘッダーより後のデータ行(届出年月日など)を含めると、日付欄の値を
    # 誤って「基準日」と誤認識してしまうため、ここでは含めない。
    raw_text = "\n".join(
        " ".join(str(c) for c in r if c is not None) for r in raw_rows[: header_idx + 1]
    )
    return ExtractResult(row_count=len(data_rows), header=header, rows=data_rows, raw_text=raw_text)


def extract_csv(path: str) -> ExtractResult:
    last_err = None
    for enc in ("utf-8-sig", "utf-16", "cp932", "utf-8"):
        for sep in (",", "\t", None):
            try:
                df = pd.read_csv(
                    path,
                    encoding=enc,
                    header=None,
                    dtype=str,
                    sep=sep,
                    engine="python" if sep is None else "c",
                )
                if df.shape[1] <= 1 and sep != ",":
                    # 区切り文字の推定に失敗した可能性が高いので次を試す
                    continue
                rows = df.values.tolist()
                result = _clean_table(rows)
                result.note = f"encoding={enc}, sep={sep!r}"
                return result
            except Exception as e:  # noqa: BLE001
                last_err = e
                continue
    return ExtractResult(row_count=0, note=f"CSV読み込み失敗: {last_err}")


KIND_SHEET_HINTS = {
    "旅館業": ["旅館", "ホテル", "簡易宿所", "簡易宿泊", "簡易宿泊所"],
    "民泊": ["住宅宿泊", "民泊"],
}


def extract_excel(path: str, kind_hint: str | None = None, sheet_hint: str | None = None) -> ExtractResult:
    try:
        xls = pd.ExcelFile(path)
    except Exception as e:  # noqa: BLE001
        return ExtractResult(row_count=0, note=f"Excel読み込み失敗: {e}")

    all_sheet_names_text = " ".join(xls.sheet_names)
    result = _extract_excel_pick(xls, kind_hint, sheet_hint)
    result.raw_text = (result.raw_text or "") + "\n" + all_sheet_names_text
    return result


def _extract_excel_pick(xls, kind_hint: str | None, sheet_hint: str | None) -> ExtractResult:
    def parse_sheet(sheet: str) -> ExtractResult | None:
        try:
            df = xls.parse(sheet_name=sheet, header=None, dtype=str)
        except Exception:  # noqa: BLE001
            return None
        result = _clean_table(df.values.tolist())
        result.sheet_name = sheet
        return result

    # 1) 明示的にシート名が指定されていればそれを使う(カンマ区切りで複数指定可)
    if sheet_hint:
        wanted = [s.strip() for s in sheet_hint.split(",") if s.strip()]
        matched = [s for s in xls.sheet_names if s in wanted]
        if matched:
            combined_rows: list[list] = []
            header: list[str] = []
            for s in matched:
                r = parse_sheet(s)
                if r is None or r.row_count == 0:
                    continue
                if not header:
                    header = r.header
                combined_rows.extend(r.rows)
            if combined_rows:
                return ExtractResult(
                    row_count=len(combined_rows), header=header, rows=combined_rows,
                    sheet_name=",".join(matched), note="指定シートを使用",
                )

    # 2) 種別(民泊/旅館業)に対応するキーワードをシート名から探す
    hints = KIND_SHEET_HINTS.get(kind_hint or "", [])
    if hints:
        matched = [s for s in xls.sheet_names if any(h in s for h in hints)]
        if matched:
            combined_rows = []
            header = []
            for s in matched:
                r = parse_sheet(s)
                if r is None or r.row_count == 0:
                    continue
                if not header:
                    header = r.header
                combined_rows.extend(r.rows)
            if combined_rows:
                return ExtractResult(
                    row_count=len(combined_rows), header=header, rows=combined_rows,
                    sheet_name=",".join(matched), note=f"種別キーワード一致シートを使用: {matched}",
                )

    # 3) どちらも手がかりがなければ、単一シートのブックは素直にそれを、
    #    複数シートあるブックは最大行数のシートを使う(ただし要確認扱いにする)
    if len(xls.sheet_names) == 1:
        r = parse_sheet(xls.sheet_names[0])
        return r or ExtractResult(row_count=0, note="読み取れるシートがありませんでした")

    best: ExtractResult | None = None
    for sheet in xls.sheet_names:
        r = parse_sheet(sheet)
        if r is None:
            continue
        if best is None or r.row_count > best.row_count:
            best = r
    if best is None:
        return ExtractResult(row_count=0, note="読み取れるシートがありませんでした")
    best.note = f"複数シートのブックでキーワード一致なし。最大行数のシート'{best.sheet_name}'を暫定使用(要確認)"
    return best


def extract_pdf(path: str) -> ExtractResult:
    import pdfplumber

    all_rows: list[list] = []
    text_fallback_pages = 0
    first_pages_text = ""
    try:
        with pdfplumber.open(path) as pdf:
            # 日付の手がかりは書類冒頭のタイトル部分にしか出てこないことが多く、
            # ページ全体のテキストを使うと本文中のデータ(届出年月日など)を
            # 誤って拾ってしまうため、1ページ目の先頭だけに絞る
            first_page_text = (pdf.pages[0].extract_text() or "") if pdf.pages else ""
            first_pages_text = first_page_text[:400]
            for page in pdf.pages:
                tables = page.extract_tables()
                line_based_rows = sum(len(t) for t in tables) if tables else 0

                # 罫線ベースの検出が「ほぼ何も取れていない」(0行、または
                # ページ本文の行数に比べて明らかに少ない=飾り枠だけ拾った)
                # ときだけ、文字位置ベースの検出を試す。両方それなりに
                # 行があるときは、通常より正確な罫線ベースを優先する。
                need_fallback = line_based_rows == 0
                if not need_fallback and line_based_rows <= 2:
                    page_text = page.extract_text() or ""
                    approx_lines = len([l for l in page_text.split("\n") if l.strip()])
                    if approx_lines > line_based_rows * 3:
                        need_fallback = True

                if need_fallback:
                    try:
                        text_table = page.extract_table(
                            table_settings={
                                "vertical_strategy": "text",
                                "horizontal_strategy": "text",
                            }
                        )
                    except Exception:  # noqa: BLE001
                        text_table = None
                    if text_table and len(text_table) > line_based_rows:
                        all_rows.extend(text_table)
                        text_fallback_pages += 1
                        continue

                for table in tables:
                    all_rows.extend(table)
    except Exception as e:  # noqa: BLE001
        return ExtractResult(row_count=0, note=f"PDF読み込み失敗: {e}")

    if not all_rows:
        return ExtractResult(row_count=0, note="PDFから表を検出できませんでした(画像PDFの可能性)")
    all_rows = [[_fix_pdf_text(c) for c in row] for row in all_rows]
    first_pages_text = _fix_pdf_text(first_pages_text)
    result = _clean_table(all_rows)
    note = f"pdfplumberで{len(all_rows)}行検出"
    if text_fallback_pages:
        note += f"(罫線なしページ{text_fallback_pages}件はテキスト位置ベースで検出)"
    result.note = note
    result.raw_text = first_pages_text + "\n" + result.raw_text
    return result


def extract_docx(path: str) -> ExtractResult:
    import docx

    try:
        doc = docx.Document(path)
    except Exception as e:  # noqa: BLE001
        return ExtractResult(row_count=0, note=f"Word読み込み失敗: {e}")

    all_rows: list[list] = []
    for table in doc.tables:
        for row in table.rows:
            all_rows.append([cell.text for cell in row.cells])

    if not all_rows:
        return ExtractResult(row_count=0, note="Word文書内に表が見つかりませんでした")
    result = _clean_table(all_rows)
    para_text = "\n".join(p.text for p in doc.paragraphs)[:400]
    result.raw_text = para_text + "\n" + result.raw_text
    return result


KANI_SHUKUSHO_HINTS = ["種別", "区分", "営業形態"]


def split_kani_shukusho(header: list[str], rows: list[list]) -> tuple[list[list], list[list]] | None:
    """旅館業の一覧を「簡易宿所」とそれ以外(旅館・ホテル等)に分ける。
    分類できる列が見つからない場合は None を返す。
    """
    col_idx = None
    for i, h in enumerate(header):
        if any(hint in h for hint in KANI_SHUKUSHO_HINTS):
            col_idx = i
            break
    if col_idx is None:
        return None

    kani_rows = []
    other_rows = []
    found_kani = False
    for r in rows:
        cell = str(r[col_idx]) if col_idx < len(r) and r[col_idx] is not None else ""
        if "簡易宿所" in cell:
            kani_rows.append(r)
            found_kani = True
        else:
            other_rows.append(r)

    if not found_kani:
        return None
    return kani_rows, other_rows


MINPAKU_RYOKAN_COL_HINTS = ["営業種別", "業種別", "施設区分", "種別", "区分"]
MINPAKU_KEYWORDS = ["住宅宿泊事業", "特区民泊", "民泊"]
RYOKAN_KEYWORDS = ["旅館", "ホテル", "簡易宿所"]


def split_minpaku_ryokan(header: list[str], rows: list[list]) -> tuple[list[list], list[list]] | None:
    """民泊(住宅宿泊事業・特区民泊)と旅館業(旅館・ホテル・簡易宿所)が
    1つの表に混在している資料を、種別列の値で2つに分ける。
    両方の種別が実際に見つかった場合のみ (民泊行, 旅館業行) を返し、
    そうでなければ(=混在資料ではない) None を返す。
    """
    col_idx = None
    for i, h in enumerate(header):
        if any(hint in h for hint in MINPAKU_RYOKAN_COL_HINTS):
            col_idx = i
            break
    if col_idx is None:
        return None

    minpaku_rows = []
    ryokan_rows = []
    for r in rows:
        cell = str(r[col_idx]) if col_idx < len(r) and r[col_idx] is not None else ""
        if any(k in cell for k in MINPAKU_KEYWORDS):
            minpaku_rows.append(r)
        elif any(k in cell for k in RYOKAN_KEYWORDS):
            ryokan_rows.append(r)

    if not minpaku_rows or not ryokan_rows:
        return None
    return minpaku_rows, ryokan_rows


def extract(path: str, fmt: str, kind_hint: str | None = None, sheet_hint: str | None = None) -> ExtractResult:
    if fmt in ("xlsx", "xls"):
        return extract_excel(path, kind_hint=kind_hint, sheet_hint=sheet_hint)
    if fmt == "csv":
        return extract_csv(path)
    if fmt == "pdf":
        return extract_pdf(path)
    if fmt == "docx":
        return extract_docx(path)
    return ExtractResult(row_count=0, note=f"未対応の形式: {fmt}")
